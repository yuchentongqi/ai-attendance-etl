import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from attendance_api import read_real_employees, get_mock_attendance_records, summarize_attendance

def generate_excel_report():
    today = datetime.now().strftime("%Y-%m-%d")
    employees = read_real_employees()
    mock_records = get_mock_attendance_records(employees, work_date=today)
    summary = summarize_attendance(mock_records)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    headers = ["部门名称", "部门ID", "考勤人员姓名", "考勤人员id", "工种", "备注", "当天上班打卡时间", "当天下班打卡时间"]
    
    ws.merge_cells('A1:H1')
    ws['A1'] = f"马来项目施工人员考勤数据表 ({today})"
    ws['A1'].font = Font(name='Microsoft YaHei', size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 25

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = Font(name='Microsoft YaHei', size=12, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18

    row_idx = 3
    dept_id_map = {} 
    dept_counter = 1000

    for emp in employees:
        uid = emp["userid"]
        dept_name = emp["dept_name"]
        if dept_name not in dept_id_map:
            dept_id_map[dept_name] = dept_counter
            dept_counter += 1
        dept_id = dept_id_map[dept_name]

        info = summary.get(uid, {})
        clock_in = info.get("clock_in", "")
        clock_out = info.get("clock_out", "")
        status = info.get("status", [])
        remark = "、".join(status) if status else "无备注信息"

        row_data = [dept_name, dept_id, emp["name"], uid, emp["work_type"], remark, clock_in, clock_out]
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Microsoft YaHei', size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row_idx].height = 18
        row_idx += 1

    for col in ws.columns:
        max_length = 0
        column_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                length = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                if length > max_length: max_length = length
        ws.column_dimensions[column_letter].width = max_length + 4

    wb.save("考勤报表.xlsx")
    print("✅ Excel 报表已生成：考勤报表.xlsx")

if __name__ == "__main__":
    generate_excel_report()