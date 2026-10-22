import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict

def analyze_attendance():
    wb = openpyxl.load_workbook("考勤报表.xlsx", data_only=True)
    ws = wb["sheet1"]
    dept_stats = defaultdict(lambda: defaultdict(int))

    for row in ws.iter_rows(min_row=3, values_only=True):
        dept_name, work_type, clock_in = row[0], row[4], row[6]
        if dept_name and work_type and clock_in:
            dept_stats[dept_name][work_type] += 1

    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    new_ws.title = "分类统计表"

    all_work_types = set()
    for types in dept_stats.values():
        all_work_types.update(types.keys())
    work_type_list = sorted(list(all_work_types))
    headers = ["部门名称"] + work_type_list + ["部门总计"]
    
    new_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = new_ws.cell(row=1, column=1, value="施工人员考勤分类统计")
    title_cell.font = Font(name='Microsoft YaHei', size=14, bold=True)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    new_ws.row_dimensions[1].height = 25

    for col_idx, header in enumerate(headers, start=1):
        cell = new_ws.cell(row=2, column=col_idx, value=header)
        cell.font = Font(name='Microsoft YaHei', size=12, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    new_ws.row_dimensions[2].height = 18

    row_idx = 3
    for dept, types in dept_stats.items():
        new_ws.cell(row=row_idx, column=1, value=dept).font = Font(name='Microsoft YaHei', size=11)
        dept_total = 0
        for col_idx, wt in enumerate(work_type_list, start=2):
            count = types.get(wt, 0)
            cell = new_ws.cell(row=row_idx, column=col_idx, value=count)
            cell.font = Font(name='Microsoft YaHei', size=11)
            cell.alignment = Alignment(horizontal='center')
            dept_total += count
        total_cell = new_ws.cell(row=row_idx, column=len(headers), value=dept_total)
        total_cell.font = Font(name='Microsoft YaHei', size=11, bold=True)
        total_cell.alignment = Alignment(horizontal='center')
        new_ws.row_dimensions[row_idx].height = 18
        row_idx += 1

    for col_idx in range(1, len(headers) + 1):
        max_length = 0
        column_letter = get_column_letter(col_idx)
        for row in range(2, row_idx):
            cell_value = new_ws.cell(row=row, column=col_idx).value
            if cell_value:
                length = sum(2 if ord(c) > 127 else 1 for c in str(cell_value))
                if length > max_length: max_length = length
        new_ws.column_dimensions[column_letter].width = max_length + 4

    new_wb.save("考勤统计表.xlsx")
    print("✅ 统计表已生成：考勤统计表.xlsx")

if __name__ == "__main__":
    analyze_attendance()